import io
from openpyxl import Workbook, load_workbook


def read_product_names(file_bytes: bytes) -> list[str]:
    wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    start_idx = 0
    col_idx = 0
    header = rows[0]
    for i, cell in enumerate(header):
        if cell and "상품명" in str(cell):
            col_idx = i
            start_idx = 1
            break

    names = []
    for row in rows[start_idx:]:
        if col_idx < len(row) and row[col_idx]:
            names.append(str(row[col_idx]).strip())
    return [n for n in names if n]


def build_result_excel(results: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "결과"
    # 기존 칸 배치는 그대로 두고(복사해 쓰는 칸이 밀리지 않게) 이름 출처/후보는 맨 뒤 별도 칸에 둔다
    ws.append(["기존 상품명", "최적화된 상품명", "최적화상품명 PC검색량", "최적화상품명 모바일검색량", "추출 키워드 (검색량순, 복사용)",
               "이름 출처", "후보1 기본 이름", "후보2 대안 이름"])

    for r in results:
        keyword_str = ",".join(k["keyword"] for k in r["keywords"])
        ws.append([
            r["original_name"],
            r["optimized_name"],
            r.get("optimized_name_pc", 0),
            r.get("optimized_name_mobile", 0),
            keyword_str,
            r.get("name_source", ""),
            r.get("base_name", ""),
            r.get("alt_name", ""),
        ])

    for col, width in zip("ABCDEFGH", [25, 25, 18, 20, 100, 22, 25, 25]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
